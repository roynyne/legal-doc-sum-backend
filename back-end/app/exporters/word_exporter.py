from docx import Document
from docx.shared import Pt
from datetime import datetime

def generate_docx(case_data: dict, output_path: str):
    doc = Document()

    # Title
    doc.add_heading("📄 Legal Case Brief", level=0)
    doc.add_paragraph(f"Case: {case_data['case_name']}")
    doc.add_paragraph(f"Generated on: {datetime.now().strftime('%d %B %Y')}")
    doc.add_paragraph("")

    # Structured Summary
    doc.add_heading("Summary", level=1)
    for section, text in case_data["structured_summary"].items():
        doc.add_heading(section.capitalize(), level=2)
        doc.add_paragraph(text)

    # Evaluation Section
    if "evaluation" in case_data:
        doc.add_page_break()
        doc.add_heading("📊 Evaluation Metrics", level=1)
        for key, val in case_data["evaluation"].items():
            if isinstance(val, float):
                val = f"{val:.2f}"
            elif isinstance(val, list):
                val = ", ".join(val)
            doc.add_paragraph(f"{key.replace('_', ' ').title()}: {val}", style='List Bullet')

    doc.save(output_path)
